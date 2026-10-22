# -*- coding: utf-8 -*-
"""Build a formatted DOCX user manual from docs/USER_MANUAL.md."""
import re, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'USER_MANUAL.md')
OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.join(ROOT, 'docs', 'USER_MANUAL.docx')
LOGO = os.path.join(ROOT, 'public', 'hgi-logo.png')

NAVY = RGBColor(0x14, 0x1A, 0x58)
INDIGO = RGBColor(0x1F, 0x2A, 0x6B)
SLATE = RGBColor(0x33, 0x41, 0x55)
GREY = RGBColor(0x64, 0x74, 0x8B)
BODY = RGBColor(0x1E, 0x29, 0x3B)
BT = chr(96)

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)

def shade_paragraph(p, fill, border=None, size=8):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)
    if border:
        pBdr = OxmlElement('w:pBdr')
        for side in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement('w:' + side)
            el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(size))
            el.set(qn('w:space'), '6'); el.set(qn('w:color'), border)
            pBdr.append(el)
        pPr.append(pBdr)

def bottom_border(p, color='D8DEEA', size=6):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    el = OxmlElement('w:bottom')
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(size))
    el.set(qn('w:space'), '4'); el.set(qn('w:color'), color)
    pBdr.append(el)
    pPr.append(pBdr)

def add_field(paragraph, instr, placeholder=''):
    r = paragraph.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = placeholder
    fld3 = OxmlElement('w:fldChar'); fld3.set(qn('w:fldCharType'), 'end')
    r._r.append(fld1); r._r.append(it); r._r.append(fld2); r._r.append(t); r._r.append(fld3)
    return r

def enable_update_fields(doc):
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true')
    settings.append(uf)

INLINE = re.compile(r'(\*\*.+?\*\*|\x60[^\x60]+\x60|\*[^*\n]+?\*)')

def add_runs(p, text, size=None, color=None, bold=False, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith(BT) and part.endswith(BT) and len(part) > 2:
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            r.font.size = Pt((size or 10.5) - 1)
            r.font.color.rgb = RGBColor(0x8B, 0x1A, 0x3A)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        if size and r.font.size is None:
            r.font.size = Pt(size)
        if color is not None and r.font.color.rgb is None:
            r.font.color.rgb = color
        if bold:
            r.bold = True
        if italic:
            r.italic = True

def style_doc(doc):
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = BODY
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.12
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:eastAsia'), 'Calibri')

    h1 = doc.styles['Heading 1']
    h1.font.name = 'Calibri'; h1.font.size = Pt(21); h1.font.bold = True
    h1.font.color.rgb = NAVY
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles['Heading 2']
    h2.font.name = 'Calibri'; h2.font.size = Pt(14); h2.font.bold = True
    h2.font.color.rgb = INDIGO
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles['Heading 3']
    h3.font.name = 'Calibri'; h3.font.size = Pt(11.5); h3.font.bold = True
    h3.font.color.rgb = SLATE
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

def setup_page(doc):
    s = doc.sections[0]
    s.page_width = Cm(21.0); s.page_height = Cm(29.7)
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.3); s.right_margin = Cm(2.3)
    s.header_distance = Cm(1.1); s.footer_distance = Cm(1.0)

def build_header_footer(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    hp.text = ''
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.4), WD_ALIGN_PARAGRAPH.RIGHT)
    r = hp.add_run('HYTech Learning Management System  ·  User Manual')
    r.font.size = Pt(8); r.font.color.rgb = GREY
    hp.add_run('\t')
    r2 = hp.add_run('Version 1.0')
    r2.font.size = Pt(8); r2.font.color.rgb = GREY
    bottom_border(hp, 'D8DEEA', 4)

    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run('HYT Global Institute   |   Page ')
    r.font.size = Pt(8); r.font.color.rgb = GREY
    add_field(fp, 'PAGE', '1')
    r = fp.add_run(' of ')
    r.font.size = Pt(8); r.font.color.rgb = GREY
    add_field(fp, 'NUMPAGES', '1')
    for run in fp.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GREY

def build_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(LOGO):
        p.add_run().add_picture(LOGO, width=Inches(1.9))

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(2)
    r = p.add_run('HYT GLOBAL INSTITUTE')
    r.font.size = Pt(11); r.bold = True; r.font.color.rgb = GREY

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run('Learning Management System')
    r.font.size = Pt(11); r.font.color.rgb = GREY

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('HYTech Learning Management System')
    r.font.size = Pt(30); r.bold = True; r.font.color.rgb = NAVY

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('User Manual')
    r.font.size = Pt(24); r.font.color.rgb = INDIGO

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('Super Admin   ·   Trainer   ·   Trainee')
    r.font.size = Pt(13); r.bold = True; r.font.color.rgb = SLATE

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run('A step-by-step guide to every screen, field, and permission in the system')
    r.font.size = Pt(10.5); r.italic = True; r.font.color.rgb = GREY

    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Cm(12.4)
    set_cell_bg(cell, '141A58')
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(6); cp.paragraph_format.space_after = Pt(6)
    r = cp.add_run('Version 1.0     |     October 2026     |     Document code HGI-LMS-UM-001')
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.bold = True

    for _ in range(6):
        doc.add_paragraph()

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    rows = [
        ('Prepared for', 'HYT Global Institute — Training Operations'),
        ('Prepared by', 'Documentation Team'),
        ('Applies to', 'HYTech LMS, build of 21 October 2026'),
        ('Status', 'Released for review'),
    ]
    for i, (k, v) in enumerate(rows):
        c0, c1 = meta.rows[i].cells
        c0.width = Cm(4.0); c1.width = Cm(8.4)
        p0 = c0.paragraphs[0]; r = p0.add_run(k)
        r.font.size = Pt(9); r.bold = True; r.font.color.rgb = GREY
        p1 = c1.paragraphs[0]; r = p1.add_run(v)
        r.font.size = Pt(9); r.font.color.rgb = SLATE
        for c in (c0, c1):
            c.paragraphs[0].paragraph_format.space_after = Pt(2)

    doc.add_page_break()

def add_callout(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    text = ' '.join(l.lstrip('>').strip() for l in lines if l.strip())
    text = re.sub(r'\s+', ' ', text)
    shade_paragraph(p, 'EEF2FB', 'C7D2FE', 12)
    add_runs(p, text, size=9.8, color=SLATE)
    return p

def add_screenshot_box(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.1)
    shade_paragraph(p, 'F1F5F9', 'CBD5E1', 10)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = INDIGO
    r2 = p.add_run('   —  placeholder, to be replaced with a real screenshot')
    r2.italic = True; r2.font.size = Pt(8.5); r2.font.color.rgb = GREY

def add_table(doc, rows):
    header = [c.strip() for c in rows[0].strip().strip('|').split('|')]
    body = []
    for line in rows[2:]:
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if len(cells) < len(header):
            cells += [''] * (len(header) - len(cells))
        body.append(cells[:len(header)])
    table = doc.add_table(rows=1, cols=len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_cell_margins(table)
    tblPr = table._tbl.tblPr
    tblW = OxmlElement('w:tblW'); tblW.set(qn('w:type'), 'pct'); tblW.set(qn('w:w'), '5000')
    tblPr.append(tblW)
    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
    for i, h in enumerate(header):
        c = hdr.cells[i]
        set_cell_bg(c, '141A58')
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(h.replace('**', ''))
        r.bold = True; r.font.size = Pt(8.3); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for idx, cells in enumerate(body):
        row = table.add_row()
        for i, val in enumerate(cells):
            c = row.cells[i]
            if idx % 2 == 1:
                set_cell_bg(c, 'F6F8FC')
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, val, size=8.3, color=BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run('Right-click the field below and choose "Update Field" if the entries are not shown.')
    r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = GREY
    p2 = doc.add_paragraph()
    add_field(p2, 'TOC \\o "1-3" \\h \\z \\u', 'Table of contents will be generated here.')
    doc.add_page_break()

def main():
    with open(SRC, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    style_doc(doc); setup_page(doc); build_header_footer(doc); build_cover(doc)

    started = False
    in_toc = False
    i = 0
    first_section = True
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if not started:
            if stripped.startswith('## Document Information'):
                started = True
            else:
                i += 1
                continue

        if in_toc:
            if stripped.startswith('## 1. Introduction'):
                in_toc = False
            else:
                i += 1
                continue

        if stripped in ('', '---'):
            i += 1
            continue
        if stripped.startswith('<div') or stripped.startswith('</div>'):
            i += 1
            continue

        if stripped.startswith('# '):
            i += 1
            continue

        if stripped.startswith('#### '):
            p = doc.add_paragraph(style='Heading 3')
            add_runs(p, stripped[5:])
            i += 1; continue

        if stripped.startswith('### '):
            p = doc.add_paragraph(style='Heading 2')
            add_runs(p, stripped[4:])
            i += 1; continue

        if stripped.startswith('## '):
            title = stripped[3:]
            p = doc.add_paragraph(style='Heading 1')
            if not first_section:
                p.paragraph_format.page_break_before = True
            first_section = False
            add_runs(p, title)
            bottom_border(p, 'C7D2FE', 8)
            if title.startswith('Table of Contents'):
                in_toc = True
                add_toc(doc)
            i += 1; continue

        if stripped.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i]); i += 1
            if len(block) >= 2:
                add_table(doc, block)
            continue

        if stripped.startswith('>'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                block.append(lines[i]); i += 1
            add_callout(doc, block)
            continue

        if stripped.startswith('[Insert Screenshot:'):
            add_screenshot_box(doc, stripped)
            i += 1; continue

        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.85)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.85))
            r = p.add_run(m.group(1) + '.\t')
            r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = NAVY
            add_runs(p, m.group(2))
            i += 1; continue

        if re.match(r'^\s{2,}-\s+', raw):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.55)
            p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run('–  '); r.font.color.rgb = GREY
            add_runs(p, re.sub(r'^\s*-\s+', '', raw))
            i += 1; continue

        if stripped.startswith('- '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run('•  '); r.font.color.rgb = NAVY; r.bold = True
            add_runs(p, stripped[2:])
            i += 1; continue

        if stripped.startswith('**Caption:**') or stripped.startswith('**What to look for:**'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.1)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, stripped, size=9.2, color=SLATE)
            i += 1; continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    enable_update_fields(doc)
    doc.save(OUT)
    print('saved', OUT, os.path.getsize(OUT))

main()
